import io
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.services import audit
from app.services.ingestion import build_chunks, extract_sections

SAMPLES = Path(__file__).resolve().parents[1] / "data" / "sample_kb"


def test_markdown_is_chunked_with_source_metadata():
    data = (SAMPLES / "company_it_handbook.md").read_bytes()
    chunks = build_chunks("company_it_handbook.md", data, {"origin": "sample"})
    assert chunks
    assert all(c.metadata["source"] == "company_it_handbook.md" for c in chunks)
    assert all(len(c.text) <= 900 for c in chunks)
    assert any("VPN" in c.text for c in chunks)


def test_pdf_pages_are_extracted_with_page_numbers():
    pdf = SAMPLES / "novacore_company_information.pdf"
    chunks = build_chunks(pdf.name, pdf.read_bytes())
    assert chunks
    assert all(isinstance(c.metadata.get("page"), int) for c in chunks)


def test_docx_paragraphs_and_tables_are_extracted():
    doc = DocxDocument()
    doc.add_paragraph("Printer setup guide")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Floor 3"
    table.rows[0].cells[1].text = "HP LaserJet"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_sections("guide.docx", buf.getvalue())[0].text
    assert "Printer setup guide" in text and "Floor 3 | HP LaserJet" in text


def test_unsupported_type_is_rejected():
    with pytest.raises(ValueError):
        extract_sections("malware.exe", b"MZ")


def test_txt_with_bad_bytes_does_not_crash():
    chunks = build_chunks("notes.txt", b"Reset password \xff\xfe via portal")
    assert "Reset password" in chunks[0].text


def test_record_and_read_back_query():
    qid = audit.record_query(
        {"question": "vpn?", "source_used": "private_kb", "latency_ms": 1200, "provider": "groq",
         "citations": [{"title": "a"}], "trace": [{"step": "router"}], "answer": "x" * 500}
    )
    listed = audit.recent_queries(10)
    assert listed[0]["id"] == qid and len(listed[0]["answer"]) == 280
    full = audit.get_query(qid)
    assert len(full["answer"]) == 500 and full["citations"] == [{"title": "a"}]


def test_stats_grounded_rate_ignores_direct_and_errors():
    for source in ["private_kb", "private_kb", "web_search", "insufficient_evidence", "direct"]:
        audit.record_query({"question": "q", "source_used": source, "latency_ms": 1000, "provider": "groq"})
    audit.record_query({"question": "q", "status": "error"})
    stats = audit.stats(7)
    assert stats["total_queries"] == 6
    assert stats["kb_answers"] == 2 and stats["web_answers"] == 1
    assert stats["grounded_rate"] == 75.0
    assert stats["errors"] == 1
    assert len(stats["trend"]) == 7
    assert stats["trend"][-1]["kb"] == 2
    assert stats["providers"] == [{"provider": "groq", "n": 5}]


def test_filters_and_search():
    audit.record_query({"question": "VPN drops", "source_used": "private_kb"})
    audit.record_query({"question": "Teams outage", "source_used": "web_search"})
    assert [q["question"] for q in audit.recent_queries(10, source="web_search")] == ["Teams outage"]
    assert [q["question"] for q in audit.recent_queries(10, search="vpn")] == ["VPN drops"]


def test_public_document_hides_server_path(tmp_path):
    f = tmp_path / "a.md"
    f.write_text("x")
    doc = {"id": "1", "name": "a.md", "path": str(f)}
    public = audit.public_document(doc)
    assert "path" not in public and public["downloadable"] is True
    assert audit.public_document({"id": "2", "name": "b", "path": ""})["downloadable"] is False
