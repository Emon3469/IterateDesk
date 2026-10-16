import io
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.core.config import get_settings
from app.rag.vectorstore import Chunk

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


@dataclass
class Section:
    text: str
    page: int | None = None


def extract_sections(name: str, data: bytes) -> list[Section]:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return [
            Section(text=(page.extract_text() or "").strip(), page=i + 1)
            for i, page in enumerate(reader.pages)
            if (page.extract_text() or "").strip()
        ]
    if suffix in {".txt", ".md"}:
        return [Section(text=data.decode("utf-8", errors="ignore"))]
    if suffix == ".docx":
        doc = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return [Section(text="\n".join(parts))]
    raise ValueError(f"Unsupported file type: {suffix}")


def build_chunks(name: str, data: bytes, extra: dict | None = None) -> list[Chunk]:
    settings = get_settings()
    splitter = RecursiveCharacterTextSplitter(chunk_size=settings.chunk_size, chunk_overlap=settings.chunk_overlap)
    chunks: list[Chunk] = []
    for section in extract_sections(name, data):
        for piece in splitter.split_text(section.text):
            if not piece.strip():
                continue
            meta = {"source": name, **(extra or {})}
            if section.page is not None:
                meta["page"] = section.page
            chunks.append(Chunk(text=piece.strip(), metadata=meta))
    return chunks
